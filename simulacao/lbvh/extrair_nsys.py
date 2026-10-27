import subprocess
import sqlite3
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


NSYS_DEFAULT = (
    "/usr/local/NVIDIA-Nsight-Compute-2026.3/"
    "host/target-linux-x64/nsys"
)


def find_nsys():
    nsys = shutil.which("nsys")
    if nsys:
        return nsys
    if Path(NSYS_DEFAULT).is_file():
        return NSYS_DEFAULT
    return None


def find_report():
    reports = sorted(Path(".").glob("*.nsys-rep"))
    if not reports:
        return None
    if len(reports) > 1:
        print("Mais de um .nsys-rep encontrado:")
        for report in reports:
            print(f"  - {report}")
        print("\nUsando o primeiro.")
    return reports[0]


def convert_to_sqlite(report, nsys):
    sqlite_file = report.with_suffix(".sqlite")
    if sqlite_file.exists():
        print(f"SQLite já existe: {sqlite_file}")
        return sqlite_file

    print(f"Convertendo para SQLite: {sqlite_file}")
    result = subprocess.run(
        [
            nsys,
            "export",
            "--type=sqlite",
            "--output",
            str(sqlite_file),
            str(report),
        ],
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        print("\nErro ao executar o Nsight Systems:")
        print(result.stderr)
        return None
    print("Conversão concluída.")
    return sqlite_file


def get_tables(connection):
    rows = connection.execute(
        """
        SELECT name
        FROM sqlite_master
        WHERE type = 'table'
        ORDER BY name;
        """
    ).fetchall()
    return {row[0] for row in rows}


def table_exists(connection, table):
    return table in get_tables(connection)


def get_columns(connection, table):
    rows = connection.execute(
        f'PRAGMA table_info("{table}")'
    ).fetchall()
    return [row[1] for row in rows]


def write_sheet(workbook, name, headers, rows):
    worksheet = workbook.create_sheet(name)
    worksheet.append(headers)
    for row in rows:
        worksheet.append(list(row))
    worksheet.freeze_panes = "A2"

    if headers and worksheet.max_row > 1:
        last_column = get_column_letter(len(headers))
        last_row = worksheet.max_row
        worksheet.auto_filter.ref = f"A1:{last_column}{last_row}"

    for column in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 60)


def export_metric_catalog(connection, workbook):
    if not table_exists(connection, "TARGET_INFO_GPU_METRICS"):
        print("[AVISO] Catálogo GPU Metrics não encontrado.")
        return

    rows = connection.execute(
        """
        SELECT DISTINCT typeName, metricName, metricId, sourceId, typeId
        FROM TARGET_INFO_GPU_METRICS
        ORDER BY typeName, metricName;
        """
    ).fetchall()
    write_sheet(
        workbook,
        "Metric_Catalog",
        ["type_name", "metric_name", "metric_id", "source_id", "type_id"],
        rows,
    )
    print(f"Catálogo de métricas: {len(rows)} entradas")


def export_gpu_metrics(connection, workbook):
    tables = get_tables(connection)
    if "GPU_METRICS" not in tables:
        print("[AVISO] GPU_METRICS não encontrada.")
        return

    if "TARGET_INFO_GPU_METRICS" in tables:
        query = """
            SELECT
                g.timestamp,
                g.rawTimestamp,
                g.typeId,
                g.metricId,
                m.typeName,
                m.metricName,
                g.value
            FROM GPU_METRICS g
            LEFT JOIN (
                SELECT DISTINCT typeId, metricId, typeName, metricName
                FROM TARGET_INFO_GPU_METRICS
            ) m ON m.typeId = g.typeId AND m.metricId = g.metricId
            ORDER BY g.timestamp;
        """
    else:
        query = """
            SELECT
                timestamp,
                rawTimestamp,
                typeId,
                metricId,
                NULL AS typeName,
                NULL AS metricName,
                value
            FROM GPU_METRICS
            ORDER BY timestamp;
        """
    rows = connection.execute(query).fetchall()
    write_sheet(
        workbook,
        "GPU_Metrics",
        ["timestamp", "raw_timestamp", "type_id", "metric_id", "type_name", "metric_name", "value"],
        rows,
    )
    print(f"Medições GPU exportadas: {len(rows)}")


def export_nvml_power_temperature(connection, workbook):
    tables = get_tables(connection)
    if not {"GENERIC_EVENTS", "GENERIC_EVENT_DATA", "GENERIC_EVENT_TYPE_FIELDS", "StringIds"}.issubset(tables):
        print("[AVISO] Tabelas genéricas do NVML não encontradas.")
        return

    query = """
        SELECT
            ge.timestamp,
            ge.rawTimestamp,
            s.value AS field_name,
            COALESCE(ged.uintVal, ged.intVal, ged.floatVal, ged.doubleVal) AS val,
            ge.globalTid
        FROM GENERIC_EVENTS ge
        JOIN GENERIC_EVENT_DATA ged ON ge.genericEventId = ged.genericEventId
        JOIN GENERIC_EVENT_TYPE_FIELDS getf ON ge.typeId = getf.typeId AND ged.fieldIdx = getf.fieldIdx
        JOIN StringIds s ON getf.fieldNameId = s.id
        WHERE s.value IN ('Power', 'Temperature')
        ORDER BY ge.timestamp;
    """
    rows = connection.execute(query).fetchall()
    if not rows:
        print("[AVISO] Nenhum registro de Power ou Temperature encontrado.")
        return

    write_sheet(
        workbook,
        "NVML_Power_Temp",
        ["timestamp", "raw_timestamp", "metric_name", "value", "global_tid"],
        rows,
    )
    print(f"Métricas NVML (Power/Temp) exportadas: {len(rows)}")


def export_vulkan_memory(connection, workbook):
    """Extrai operações de alocação e gerenciamento de memória do Vulkan (Item 14)"""
    if not table_exists(connection, "VULKAN_MEMORY_OPERATION"):
        print("[AVISO] VULKAN_MEMORY_OPERATION não encontrada.")
        return

    rows = connection.execute(
        """
        SELECT
            id,
            gpu,
            rangeStart,
            rangeEnd,
            CASE WHEN rangeEnd IS NOT NULL AND rangeStart IS NOT NULL
                 THEN (rangeEnd - rangeStart) / 1000000.0 ELSE NULL END AS duration_ms,
            memorySize,
            memoryOffset,
            contextId,
            objectHandle
        FROM VULKAN_MEMORY_OPERATION
        ORDER BY rangeStart;
        """
    ).fetchall()

    write_sheet(
        workbook,
        "Vulkan_Memory",
        [
            "id",
            "gpu",
            "range_start",
            "range_end",
            "duration_ms",
            "memory_size_bytes",
            "memory_offset",
            "context_id",
            "object_handle",
        ],
        rows,
    )
    print(f"Operações de memória Vulkan exportadas: {len(rows)}")


def export_nvtx(connection, workbook):
    if not table_exists(connection, "NVTX_EVENTS"):
        print("[AVISO] NVTX_EVENTS não encontrada.")
        return

    rows = connection.execute(
        """
        SELECT
            start,
            end,
            CASE WHEN end IS NOT NULL THEN (end - start) / 1000000.0 ELSE NULL END AS duration_ms,
            eventType,
            rangeId,
            category,
            text,
            globalTid,
            endGlobalTid,
            jsonText
        FROM NVTX_EVENTS
        ORDER BY start;
        """
    ).fetchall()
    write_sheet(
        workbook,
        "NVTX",
        ["start", "end", "duration_ms", "event_type", "range_id", "category", "text", "global_tid", "end_global_tid", "json_text"],
        rows,
    )
    print(f"Eventos NVTX exportados: {len(rows)}")


def export_vulkan_api(connection, workbook):
    if not table_exists(connection, "VULKAN_API"):
        print("[AVISO] VULKAN_API não encontrada.")
        return

    rows = connection.execute(
        """
        SELECT
            v.id,
            v.start,
            v.end,
            CASE WHEN v.end IS NOT NULL THEN (v.end - v.start) / 1000000.0 ELSE NULL END AS duration_ms,
            s.value AS function_name,
            v.globalTid,
            v.correlationId,
            v.endCorrelationId,
            v.contextId,
            v.objectNameId
        FROM VULKAN_API v
        LEFT JOIN StringIds s ON s.id = v.nameId
        ORDER BY v.start;
        """
    ).fetchall()
    write_sheet(
        workbook,
        "Vulkan_API",
        ["id", "start", "end", "duration_ms", "function", "global_tid", "correlation_id", "end_correlation_id", "context_id", "object_name_id"],
        rows,
    )
    print(f"Chamadas Vulkan exportadas: {len(rows)}")


def export_context_switches(connection, workbook):
    if not table_exists(connection, "GPU_CONTEXT_SWITCH_EVENTS"):
        print("[AVISO] GPU_CONTEXT_SWITCH_EVENTS não encontrada.")
        return

    rows = connection.execute(
        """
        SELECT timestamp, contextId, globalPid, gpuId, vmId, seqNo, tag
        FROM GPU_CONTEXT_SWITCH_EVENTS
        ORDER BY timestamp;
        """
    ).fetchall()
    write_sheet(
        workbook,
        "Context_Switches",
        ["timestamp", "context_id", "global_pid", "gpu_id", "vm_id", "sequence_number", "tag"],
        rows,
    )
    print(f"Eventos de troca de contexto: {len(rows)}")


def export_gpu_info(connection, workbook):
    if not table_exists(connection, "TARGET_INFO_GPU"):
        return
    headers = get_columns(connection, "TARGET_INFO_GPU")
    rows = connection.execute("SELECT * FROM TARGET_INFO_GPU;").fetchall()
    write_sheet(workbook, "GPU_Info", headers, rows)


def export_metadata(connection, workbook):
    tables = get_tables(connection)
    rows = []
    for table in ("META_DATA_CAPTURE", "META_DATA_EXPORT"):
        if table not in tables:
            continue
        data = connection.execute(f'SELECT name, value FROM "{table}" ORDER BY name;').fetchall()
        for name, value in data:
            rows.append((table, name, value))
    if rows:
        write_sheet(workbook, "Metadata", ["source_table", "name", "value"], rows)


def export_database_structure(connection, workbook):
    rows = []
    for table in sorted(get_tables(connection)):
        for column in get_columns(connection, table):
            rows.append((table, column))
    write_sheet(workbook, "Database_Structure", ["table", "column"], rows)


def main():
    report = find_report()
    if report is None:
        print("Nenhum arquivo .nsys-rep encontrado.")
        return

    print(f"Arquivo encontrado: {report}")
    nsys = find_nsys()
    if nsys is None:
        print("Erro: executável 'nsys' não encontrado.")
        return

    sqlite_file = convert_to_sqlite(report, nsys)
    if sqlite_file is None:
        return

    excel_file = report.with_name(report.stem + "_dados.xlsx")
    print(f"\nGerando Excel: {excel_file}")

    connection = sqlite3.connect(sqlite_file)
    try:
        workbook = Workbook()
        workbook.remove(workbook.active)

        export_metric_catalog(connection, workbook)
        export_gpu_metrics(connection, workbook)
        export_nvml_power_temperature(connection, workbook)
        export_vulkan_memory(connection, workbook)  # <--- Chamada adicionada aqui!
        export_nvtx(connection, workbook)
        export_vulkan_api(connection, workbook)
        export_context_switches(connection, workbook)
        export_gpu_info(connection, workbook)
        export_metadata(connection, workbook)
        export_database_structure(connection, workbook)

        if not workbook.sheetnames:
            print("[ERRO] Nenhuma planilha foi gerada.")
            return

        workbook.save(excel_file)
    finally:
        connection.close()

    print(f"\nExcel criado com sucesso: {excel_file}")


if __name__ == "__main__":
    main()
